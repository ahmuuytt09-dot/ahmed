#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""معاينة بصرية: يرسم ورقة Excel إلى صورة PNG للتحقق من التصميم."""
import datetime
import re
import sys

import openpyxl
from PIL import Image, ImageDraw, ImageFont

REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
BLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
PX = 7.2          # بكسل لكل وحدة عرض عمود
RPX = 1.34        # بكسل لكل وحدة ارتفاع صف
START = datetime.date(2026, 6, 1)


def fnt(sz, bold=False):
    return ImageFont.truetype(BLD if bold else REG, max(8, int(sz * 1.30)))


def argb(c):
    if c is None:
        return None
    v = getattr(c, "rgb", None)
    if not isinstance(v, str):
        return None
    v = v[-6:]
    return tuple(int(v[i:i + 2], 16) for i in (0, 2, 4)) if len(v) == 6 else None


def shown(cell):
    """قيمة تقريبية للعرض (الملف لا يحوي نتائج محسوبة)."""
    v = cell.value
    if not isinstance(v, str) or not v.startswith("="):
        if isinstance(v, (int, float)) and cell.number_format.startswith("[$-409]d"):
            return START.strftime("%a %d-%b")
        return "" if v is None else str(v)
    f = v
    m = re.match(r"=Dashboard!\$[A-Z]+\$\d+\+(\d+)$", f)
    if m:
        return (START + datetime.timedelta(days=int(m.group(1)))).strftime("%a %d-%b")
    if f.startswith('="WEEK '):
        return "WEEK 01   \u2022   01 Jun 2026  \u2013  07 Jun 2026"
    if "Reporting year" in f:
        return "Environment, Health & Safety  \u2022  Reporting year Jun 2026 \u2013 May 2027"
    if f.startswith("=DATE(2026,6,1)+("):
        return START.strftime("%A, %d %B %Y")
    if re.match(r"=[A-Z]+\d+\+6$", f):
        return (START + datetime.timedelta(days=6)).strftime("%A, %d %B %Y")
    if "0.00" in cell.number_format or cell.number_format == "0.00":
        return "0.00"
    return "0"


def render(path, sheet, out, max_row=None, max_col=None):
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    mr = max_row or ws.max_row
    mc = max_col or ws.max_column

    xs, x = [0], 0
    for c in range(1, mc + 1):
        w = ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 8.43
        x += int(w * PX)
        xs.append(x)
    ys, y = [0], 0
    for r in range(1, mr + 1):
        h = ws.row_dimensions[r].height or 15
        y += int(h * RPX)
        ys.append(y)

    img = Image.new("RGB", (xs[-1] + 2, ys[-1] + 2), "white")
    d = ImageDraw.Draw(img)

    merged = {}
    for rng in ws.merged_cells.ranges:
        for rr in range(rng.min_row, rng.max_row + 1):
            for cc in range(rng.min_col, rng.max_col + 1):
                merged[(rr, cc)] = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)

    drawn = set()
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            if (r, c) in drawn:
                continue
            cell = ws.cell(r, c)
            r1, c1, r2, c2 = merged.get((r, c), (r, c, r, c))
            if (r1, c1) != (r, c):
                continue
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1):
                    drawn.add((rr, cc))
            if r2 > mr or c2 > mc:
                r2, c2 = min(r2, mr), min(c2, mc)
            x0, y0, x1, y1 = xs[c1 - 1], ys[r1 - 1], xs[c2], ys[r2]

            src = ws.cell(r1, c1)
            fill = argb(src.fill.fgColor) if src.fill and src.fill.patternType else None
            if fill:
                d.rectangle([x0, y0, x1, y1], fill=fill)
            if src.border:
                for side, pts in (("left", [x0, y0, x0, y1]), ("right", [x1, y0, x1, y1]),
                                  ("top", [x0, y0, x1, y0]), ("bottom", [x0, y1, x1, y1])):
                    s = getattr(src.border, side)
                    if s and s.style:
                        d.line(pts, fill=argb(s.color) or (150, 150, 150),
                               width=2 if s.style == "medium" else 1)

            txt = shown(src)
            if not txt:
                continue
            f = src.font
            fo = fnt(f.size or 11, bool(f.bold))
            col = argb(f.color) or (0, 0, 0)
            al = src.alignment.horizontal or ("left" if isinstance(src.value, str) else "right")
            lines = txt.split("\n")
            th = sum(fo.getbbox(l)[3] - fo.getbbox(l)[1] + 4 for l in lines)
            ty = y0 + max(2, ((y1 - y0) - th) // 2)
            for ln in lines:
                bb = d.textbbox((0, 0), ln, font=fo)
                w = bb[2] - bb[0]
                if al == "center":
                    tx = x0 + ((x1 - x0) - w) // 2
                elif al == "right":
                    tx = x1 - w - 6
                else:
                    tx = x0 + 6
                d.text((max(tx, x0 + 2), ty), ln, font=fo, fill=col)
                ty += bb[3] - bb[1] + 4

    img.save(out)
    print("rendered", out, img.size)


if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2], sys.argv[3],
           int(sys.argv[4]) if len(sys.argv) > 4 else None,
           int(sys.argv[5]) if len(sys.argv) > 5 else None)
