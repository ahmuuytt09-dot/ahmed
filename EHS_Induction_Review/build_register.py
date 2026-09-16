"""Build the KAZ EHS induction audit register workbook + Word version of the revised induction."""
import re, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ---------- parse report tables ----------
md = open('KAZ_HSSE_Induction_Review_Report_AR.md', encoding='utf-8').read()
lines = md.split('\n')

def parse_tables(lines):
    tables, i, head, pending = [], 0, '', ''
    while i < len(lines):
        ln = lines[i]
        if ln.startswith('#'):
            pending = ln.lstrip('# ').strip()
        if ln.strip().startswith('|') and i+1 < len(lines) and set(lines[i+1].replace('|','').strip()) <= set('-: '):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                if not set(lines[i].replace('|','').strip()) <= set('-: '):
                    rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            tables.append((head or pending, rows)); continue
        i += 1
    return tables

T = {h: r for h, r in parse_tables(lines)}

# ---------- workbook ----------
wb = openpyxl.Workbook()
wb.remove(wb.active)

FN = 'Segoe UI'
HDR_FILL = PatternFill('solid', fgColor='0F3D5C')
HDR_FONT = Font(name=FN, bold=True, color='FFFFFF', size=11)
CELL_FONT = Font(name=FN, size=10)
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT = PatternFill('solid', fgColor='EAF1F6')

RED = Font(name=FN, size=10, bold=True, color='C00000')
ORG = Font(name=FN, size=10, bold=True, color='BF6000')

def add_sheet(name, title, rows, widths=None, severity_col=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    ws['A1'] = title
    ws['A1'].font = Font(name=FN, bold=True, size=13, color='0F3D5C')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(rows[0])))
    ws.row_dimensions[1].height = 24
    for j, h in enumerate(rows[0], 1):
        c = ws.cell(row=3, column=j, value=h)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, row in enumerate(rows[1:], 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = CELL_FONT, BORDER
            c.alignment = Alignment(vertical='top', wrap_text=True, horizontal='right')
            if i % 2 == 1:
                c.fill = ALT
            if severity_col and j == severity_col and '🔴' in str(v):
                c.font = RED
            elif severity_col and j == severity_col and '🟠' in str(v):
                c.font = ORG
    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f'A3:{get_column_letter(len(rows[0]))}{3+len(rows)-1}'
    default_w = 28
    for j in range(1, len(rows[0])+1):
        w = (widths[j-1] if widths and j-1 < len(widths) else default_w)
        ws.column_dimensions[get_column_letter(j)].width = w
    for i in range(4, 4+len(rows)-1):
        ws.row_dimensions[i].height = None
    return ws

# 1. Errors
add_sheet('Errors-الأخطاء', 'سجل أخطاء ملف الإندكشن (Error & Correction Register) — 120 ملاحظة',
          T['5. سجل الأخطاء والتصحيحات المباشرة (Error & Correction Register)'],
          widths=[9,10,42,14,38,52,26,10], severity_col=8)

# 2. Gaps (merge 6.1..6.4)
gap_rows = [['#','المتطلب الناقص','المرجع','الأهمية لمشروع KAZ','المحتوى الموصى بإضافته','الأولوية']]
for k in ['6.1 الفجوات المتعلقة بإطار Siemens Energy وسياساتها',
          '6.2 الفجوات المتعلقة بالمخاطر التشغيلية للمشروع',
          '6.3 الفجوات المتعلقة بالتحكم اليومي والسلوك',
          '6.4 الفجوات التنظيمية والإدارية في ملف الاندكشن نفسه']:
    gap_rows += T[k][1:]
add_sheet('Gaps-الفجوات', 'تحليل الفجوات — المتطلبات الواجب إضافتها للإندكشن (45 فجوة)',
          gap_rows, widths=[9,44,28,44,54,10], severity_col=6)

# 3. Conflicts
add_sheet('Conflicts-التعارضات', 'التعارضات الحرجة التي تتطلب قراراً إدارياً (12 تعارضاً)',
          T['7. التعارضات الحرجة التي تتطلب قراراً إدارياً موثّقاً (Conflicts Requiring Decision)'],
          widths=[9,22,34,44,44,20])

# 4. Compliance matrix
add_sheet('Compliance-المطابقة', 'مصفوفة التوافق مع متطلبات Siemens Energy',
          T['8. مصفوفة التوافق: متطلبات الإندكشن في Siemens Energy vs. الوضع الحالي'],
          widths=[50,28,30,46])

# 5. CAPA + KPI
add_sheet('CAPA-خطة العمل', 'خطة العمل التصحيحية والوقائية (CAPA)',
          T['9. خطة العمل التصحيحية والوقائية (CAPA)'],
          widths=[9,46,12,12,26,18,26], severity_col=4)

# 6. Coverage analysis
add_sheet('Coverage-التغطية', 'نتائج الفحص الآلي لتغطية الموضوعات الإلزامية في الملف الحالي',
          T['3.2 نتائج الفحص الآلي لتغطية الموضوعات الإلزامية'],
          widths=[46,16,40])

# 7. Critical findings
add_sheet('Findings-النتائج الحرجة', 'النتائج الحرجة (Critical Findings F-01…F-12)',
          T['1.2 أبرز اثنتي عشرة نتيجة حرجة (Critical Findings)'],
          widths=[9,58,50,34])

# 8. Revised TOC
add_sheet('Revised-TOC-الهيكل', 'الهيكل المقترح للإندكشن Rev.01',
          T['الملحق A — الهيكل المقترح للإندكشن (Rev.01)'],
          widths=[9,58,18,20])

# 9. Approval checklist
add_sheet('Approval-قائمة الاعتماد', 'قائمة تحقق اعتماد الإندكشن قبل الاستخدام',
          T['الملحق D — قائمة تحقق اعتماد الإندكشن قبل الاستخدام (Approval Checklist)'],
          widths=[9,80,14])

# 10. Quick reference card
add_sheet('QuickCard-البطاقة المرجعية', 'بطاقة مرجعية سريعة — ماذا أفعل إذا…؟',
          T['الملحق C — بطاقة مرجعية سريعة (ماذا أفعل إذا…؟)'],
          widths=[34,86])

wb.save('KAZ_HSSE_Induction_Review_Register.xlsx')
print('workbook saved:', wb.sheetnames)

# ---------- Word version of the revised induction ----------
try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    path = 'KAZ_Management_Staff_EHS_Induction_Rev01_EN.md'
    for raw in open(path, encoding='utf-8').read().split('\n'):
        s = raw.rstrip()
        if s.startswith('### '):
            doc.add_heading(s[4:], level=2)
        elif s.startswith('## '):
            doc.add_heading(s[3:], level=1)
        elif s.startswith('# '):
            doc.add_heading(s[2:], level=0)
        elif s.startswith('|'):
            doc.add_paragraph(s)
        elif s.startswith('- '):
            doc.add_paragraph(s[2:], style='List Bullet')
        elif s.strip() == '---':
            doc.add_paragraph('')
        else:
            doc.add_paragraph(s)
    doc.save('KAZ_Management_Staff_EHS_Induction_Rev01.docx')
    print('docx saved')
except Exception as e:
    print('docx error:', e)
