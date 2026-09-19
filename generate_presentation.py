#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
import os

# Siemens Energy Branding Colors
DEEP_PETROL = RGBColor(0x0E, 0x2F, 0x3E)  # #0E2F3E Deep Teal/Petrol background
DARK_PETROL = RGBColor(0x08, 0x24, 0x30)  # darker variant
TEAL_ACCENT = RGBColor(0x00, 0x99, 0x99)  # #009999 Siemens Teal
TEAL_LIGHT = RGBColor(0x00, 0xBF, 0xBF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xE6, 0xE6, 0xE6)
LIGHT_GREY2 = RGBColor(0xCC, 0xCC, 0xCC)
AMBER = RGBColor(0xFF, 0xB9, 0x00)  # #FFB900 Warning
RED_CRITICAL = RGBColor(0xE3, 0x06, 0x13)  # Siemens Red
RED_DARK = RGBColor(0xC8, 0x10, 0x2E)
BLACK = RGBColor(0x00, 0x00, 0x00)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, fill_color=None, line_color=None):
    shape = slide.shapes.add_shape(1, left, top, width, height) # rectangle
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

def add_text_box(slide, left, top, width, height, text, font_size=14, bold=False, color=WHITE, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_para(text_frame, text, font_size=14, bold=False, color=WHITE, space_after=Pt(6), alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.space_after = space_after
    p.alignment = alignment
    return p

# SLIDE DATA
slides_data = [
    {
        "type": "title",
        "title": "SUBCONTRACTOR EHS NON-COMPLIANCE & SAFETY VIOLATION REPORT",
        "subtitle": "KAZ Power Plant Upgrade Project | Siemens Energy",
        "meta": "Classification: CONFIDENTIAL – MANAGEMENT EYES ONLY\nDate: September 2026\nPrepared by: EHS Manager – Siemens Energy\nRef: SE-EHS-NC-2026-001 | ISO 45001:2018 & ISO 14001:2015 Aligned",
        "visual": "Full-bleed Deep Petrol background (#0E2F3E). Top-right placeholder for Siemens Energy logo (White). Center-aligned title in WHITE Bold 32pt, subtitle TEAL LIGHT 18pt. Amber horizontal rule (3pt) below title. Bottom footer: Light Grey thin line with disclaimer. Use high-contrast white typography. No imagery to maintain authority."
    },
    {
        "type": "issue",
        "number": "01",
        "title": "Unauthorized Work – Execution Without Approved Permit to Work (PTW)",
        "non_compliance": [
            "• Subcontractor workforce observed executing high-risk activities (excavation & mechanical works) without a valid, approved Permit to Work.",
            "• Direct violation of Siemens Energy Instruction: EHS General Requirements for Contractors (Sec 5.3 PTW) and KAZ Project EHS Plan Sec 7.2.",
            "• No evidence of Task Risk Assessment (TRA) review, energy isolation verification, or SIMOPS coordination.",
            "• PTW logbook not maintained; permits not displayed at work front as required."
        ],
        "risk_level": "CRITICAL",
        "risk_desc": "Potential for fatality, uncontrolled release of hazardous energy, fire/explosion, and simultaneous operations (SIMOPS) conflict. Represents a total loss of control barrier – a leading indicator of major accident potential. Breach of ISO 45001 Cl. 8.1.2.",
        "corrective": [
            "1. IMMEDIATE STOP WORK order for all activities without PTW – enforce until full compliance verified.",
            "2. Mandatory re-training of all subcontractor supervision on Siemens Energy PTW System within 24 hours.",
            "3. Initiate formal NCR and contractual warning letter – Level 1.",
            "4. EHS to conduct full PTW system audit and spot-checks daily for next 14 days.",
            "5. Supervisor accountability: removal of non-compliant supervisors if recurrence observed."
        ],
        "visual": "Layout: Left vertical accent bar in RED_CRITICAL (0.2 inch). Top header bar: DEEP_PETROL with TEAL_ACCENT line. Title in WHITE 24pt Bold. 3-column structure: Left = Non-Compliance Details (WHITE on dark), Center = Risk Level badge (RED background, WHITE text CRITICAL, icon ⚠), Right = Corrective Action (LIGHT_GREY background box with DARK_PETROL text). Use icons: document with X for PTW."
    },
    {
        "type": "issue",
        "number": "02",
        "title": "Site Access & Onboarding – Personnel Without EHS Site Induction",
        "non_compliance": [
            "• Multiple subcontractor personnel found on site and actively working without completing mandatory Siemens Energy EHS Site Induction.",
            "• No induction cards/stickers verified; onboarding register incomplete and not cross-checked at gate.",
            "• Violation of Project EHS Plan Sec 6.1 – Competency & Training and local regulatory requirement.",
            "• Indicates systemic failure in access control and subcontractor self-verification process."
        ],
        "risk_level": "CRITICAL",
        "risk_desc": "Workers unaware of site-specific hazards, emergency assembly points, muster procedures, and life-saving rules. High probability of unsafe acts, delayed emergency response, and regulatory non-compliance leading to project stoppage.",
        "corrective": [
            "1. Immediate removal of all non-inducted personnel from site.",
            "2. Implement 100% gate check – Security + EHS joint verification – no induction, no entry.",
            "3. Subcontractor to submit full manpower list vs induction records within 12 hours – gap closure plan.",
            "4. Introduce biometric/QR-based induction tracking system.",
            "5. Formal warning and cost back-charge for EHS re-induction sessions."
        ],
        "visual": "Use ID badge / Access Denied icon. Risk badge RED_CRITICAL. Non-compliance box with AMBER left border (4pt) to highlight systemic failure. Add small footer note: 'Zero Tolerance: No Induction = No Work – Life Saving Rule #1'. Background DEEP_PETROL, content cards in slightly lighter #143D4E."
    },
    {
        "type": "issue",
        "number": "03",
        "title": "Initial PPE Non-Compliance – Failure to Wear Mandatory PPE",
        "non_compliance": [
            "• Direct observation: Workers operating without mandatory PPE – safety helmets not chin-strapped, safety glasses, gloves, and high-visibility vests.",
            "• Violation of Siemens Energy Life-Saving Rules and PPE Standard (EN/ANSI compliant) as per EHS Principles.",
            "• Subcontractor supervision present on site but failing to enforce.",
            "• PPE issued but not worn – behavioral and supervisory failure, not availability."
        ],
        "risk_level": "HIGH",
        "risk_desc": "Immediate risk of head injury, eye injury, hand lacerations, and struck-by incidents. Demonstrates weak safety culture and supervisory oversight. Repeat exposure significantly increases LTIFR/TRIR.",
        "corrective": [
            "1. On-the-spot correction and toolbox talk – documented with photographs.",
            "2. Issue Safety Observation Cards and formal EHS violation notice.",
            "3. Mandatory PPE compliance campaign – daily PPE checks by HSE Officers.",
            "4. Link PPE compliance to subcontractor KPI and payment milestone.",
            "5. Supervisor accountability briefing – 3 strikes removal policy communicated."
        ],
        "visual": "Grid of PPE icons (helmet, glasses, gloves) with red X overlay for missing items. Risk badge AMBER with DARK_PETROL text (HIGH). Use split layout: left text, right photo placeholder with caption 'Insert Site Photo Evidence – Fig 3.1'. Highlight box: AMBER #FFB900 for attention."
    },
    {
        "type": "issue",
        "number": "04",
        "title": "Excavation Safety – Inadequate Hard Barricades & Physical Protection",
        "non_compliance": [
            "• Excavation zones >1.2m depth observed without hard barricades; only loose warning tape used.",
            "• No edge protection, no safe access/egress (ladders), spoil pile within 1m of excavation edge.",
            "• Non-compliance with OSHA 29 CFR 1926 Subpart P and Siemens Energy Excavation Safety Standard.",
            "• Night-time visibility absent – no reflective markers or lighting."
        ],
        "risk_level": "CRITICAL",
        "risk_desc": "High potential for fatal fall into excavation, cave-in engulfment, and personnel/vehicle fall. Excavation is classified as High-Risk activity with fatality history in industry. Regulatory violation may trigger authority intervention.",
        "corrective": [
            "1. IMMEDIATE hard barricading with scaffold tubes / concrete barriers around all open excavations.",
            "2. Provide compliant access ladders every 7.5m and keep spoil 1.5m+ away.",
            "3. Excavation permit re-validation with geotechnical check.",
            "4. Install reflective delineators and illumination for low-light hours.",
            "5. Competent person inspection checklist to be submitted daily before start of work."
        ],
        "visual": "Diagram-style slide: Show cross-section of excavation with red highlights for missing barriers. Use RED_CRITICAL callout arrows. Risk badge RED. Include standard reference: 'Required: Hard Barricade per SE-GEN-EHS-007'. Photo placeholder box with red border."
    },
    {
        "type": "issue",
        "number": "05",
        "title": "Directives Non-Compliance – Failure to Follow Site EHS Instructions",
        "non_compliance": [
            "• Repeated failure to comply with verbal and written EHS instructions issued by Siemens Energy EHS team.",
            "• Toolbox Talks (TBT) not conducted, or conducted without documentation/attendance.",
            "• Previous NCRs and observation reports closed without evidence.",
            "• Demonstrates willful disregard for Project EHS Management System (ISO 45001 Cl. 7.3 & 8.1)."
        ],
        "risk_level": "HIGH",
        "risk_desc": "Erosion of EHS authority and safety culture. Creates precedent for further violations. If directives are optional, all critical controls (LOTO, Working at Height) are at risk of failure. Systemic cultural failure.",
        "corrective": [
            "1. Escalate to Project Management – formal letter to Subcontractor Top Management.",
            "2. Implement Daily EHS Directives Log – sign-off by Subcontractor PM and SE EHS Manager.",
            "3. Hold EHS Leadership stand-down meeting with all subcontractor foremen – documented commitment.",
            "4. Link non-compliance to contractual penalty clause.",
            "5. Weekly EHS compliance scorecard to be reviewed in project management meeting."
        ],
        "visual": "Timeline of ignored directives: show 3-4 dated instruction vs non-compliance icons. Use AMBER warning icons. Title bar with 'Culture of Non-Compliance' in RED. Design: White text on DEEP_PETROL, central AMBER exclamation triangle."
    },
    {
        "type": "issue",
        "number": "06",
        "title": "Site Traffic & Signage – Poor, Damaged, or Missing Safety Signage",
        "non_compliance": [
            "• Speed limit signs missing/damaged; no designated parking signage; pedestrian walkways not marked.",
            "• Vehicle-pedestrian interface unmanaged – no flagmen at high-risk intersections.",
            "• Violation of Project Traffic Management Plan and ISO 39001 Road Traffic Safety.",
            "• Observed speeding >30 km/h in construction zone (limit 15 km/h)."
        ],
        "risk_level": "HIGH",
        "risk_desc": "High risk of vehicle-pedestrian collision, equipment damage, and fatality. Poor signage indicates lack of site planning and increases liability in case of incident. Emergency vehicle access may be impeded.",
        "corrective": [
            "1. Immediate reinstatement of all traffic signs per approved Traffic Management Plan – retro-reflective grade.",
            "2. Establish dedicated pedestrian routes with hard segregation (Jersey barriers).",
            "3. Deploy trained flagmen and speed monitoring – radar gun checks.",
            "4. Conduct traffic safety campaign and driver re-induction.",
            "5. EHS + Logistics joint daily inspection of traffic controls."
        ],
        "visual": "Map-style schematic placeholder: show site roads with missing signs marked in RED. Use icons for speed limit, parking, pedestrian. Risk badge AMBER. Visual suggestion: Before/After photo placeholders side by side."
    },
    {
        "type": "issue",
        "number": "07",
        "title": "Waste Management – Lack of Segregation & Labeling",
        "non_compliance": [
            "• Waste containers observed without segregation – hazardous, recyclable, and general waste mixed.",
            "• No clear labeling, color-coding, or MSDS reference for hazardous waste bins.",
            "• Oil-contaminated rags and filters disposed in general waste – environmental violation.",
            "• Non-compliance with ISO 14001:2015 Cl. 8.1, Project Waste Management Plan, and local EPA regulations."
        ],
        "risk_level": "HIGH",
        "risk_desc": "Environmental pollution risk (soil/groundwater contamination), fire hazard from incompatible waste mixing, regulatory fine, and reputational damage to Siemens Energy. Potential for legal non-compliance.",
        "corrective": [
            "1. Immediate provision of color-coded, labeled waste skips (per SE Environmental Standard).",
            "2. Clean-up and correct segregation of existing mixed waste by competent team.",
            "3. Training on Waste Management – focus on hazardous waste handling.",
            "4. Appoint Waste Management Focal Point from subcontractor.",
            "5. Weekly waste audit – track waste volumes and disposal manifests."
        ],
        "visual": "Use 3-bin color system graphic: Green/Blue/Red bins with labels. Current state photo placeholder with RED border showing mixed waste. Environmental icon. Corrective action box with LIGHT_GREY background and green check icons."
    },
    {
        "type": "issue",
        "number": "08",
        "title": "Equipment Integrity – Active Oil & Fuel Leaks from Heavy Machinery",
        "non_compliance": [
            "• Heavy equipment (excavator, loader) operating with active hydraulic oil and fuel leaks.",
            "• Drip trays not provided; soil contamination observed under equipment parking area.",
            "• No preventive maintenance records; daily inspection checklists falsified or not done.",
            "• Violation of Equipment Safety and Environmental Protection requirements."
        ],
        "risk_level": "CRITICAL",
        "risk_desc": "Fire and explosion hazard, slip hazard, soil and groundwater contamination, equipment failure leading to loss of control/incident. Represents breakdown of maintenance management system. Breach of ISO 14001 and PUWER regulations.",
        "corrective": [
            "1. IMMEDIATE grounding of leaking equipment – tagged 'DO NOT USE'.",
            "2. Contaminated soil to be excavated and disposed as hazardous waste – spill response.",
            "3. Subcontractor to submit full equipment maintenance history and third-party inspection certificates within 48h.",
            "4. Implement mandatory daily equipment checklist – verified by SE EHS/Mechanical.",
            "5. Introduction of drip trays and spill kits at all equipment parking zones."
        ],
        "visual": "Risk badge RED_CRITICAL with oil drop icon. Layout: Left non-compliance text, Right split: Top = Photo placeholder of leak with RED circle, Bottom = Spill containment diagram. Use AMBER highlight for 'FIRE HAZARD'."
    },
    {
        "type": "issue",
        "number": "09",
        "title": "Hazardous Material Storage – Diesel Storage Without Secondary Containment",
        "non_compliance": [
            "• Diesel storage area (approx. 2000L) observed without secondary containment/bunding (110% capacity required).",
            "• No isolation distance from ignition sources; no fire extinguisher, no hazard signage, no spill kit.",
            "• Storage on bare ground – no impervious base, no earthing/bonding for fuel transfer.",
            "• Direct violation of HSG 176, EPA Spill Prevention, and SE Hazardous Materials Standard."
        ],
        "risk_level": "CRITICAL",
        "risk_desc": "Major fire/explosion risk, catastrophic environmental release, soil/groundwater contamination leading to regulatory prosecution. Potential for major accident hazard (MAH). Immediate danger to life and environment.",
        "corrective": [
            "1. STOP all fueling operations – secure area with fire watch until compliant.",
            "2. Provide compliant bunded storage – 110% secondary containment, impervious floor, roofed, ventilated.",
            "3. Install fire extinguishers (2x 9kg DCP), HAZCHEM signage, spill kit (150L), and earthing system.",
            "4. Develop and communicate Fuel Storage & Handling SOP and emergency response drill.",
            "5. Location to be approved by SE EHS – minimum 15m from excavations/buildings/drains."
        ],
        "visual": "Critical safety slide: Full-width RED_CRITICAL top banner. Show compliant vs non-compliant bund diagram side by side. Use HAZMAT diamond icons. Add '110% BUNDING REQUIRED' callout in AMBER bold. Photo placeholder with thick RED border."
    },
    {
        "type": "issue",
        "number": "10",
        "title": "Recurring Safety Violations – Persistent PPE Violations Despite Warnings",
        "non_compliance": [
            "• Continued and repeated PPE violations observed after 3 formal warnings, 5 observation cards, and 2 NCRs.",
            "• Same workgroup and supervision involved – indicates intentional non-compliance, not lack of awareness.",
            "• Trend analysis shows 60% increase in PPE violations in last 2 weeks (from Daily EHS Reports).",
            "• Subcontractor management failing to enforce disciplinary procedure as per contract."
        ],
        "risk_level": "CRITICAL",
        "risk_desc": "Normalization of deviance – precursor to serious injury/fatality. Demonstrates total failure of subcontractor safety leadership and culture. If PPE (last line of defense) is ignored, higher-level controls will also fail. Project reputation and Siemens Energy Zero Harm commitment at risk.",
        "corrective": [
            "1. Invoke contractual penalty – stop work for repeat offender crew and mandatory leadership engagement.",
            "2. Remove and replace non-compliant supervision – require CVs and EHS competency assessment.",
            "3. Implement Behavior-Based Safety (BBS) program with positive/negative reinforcement.",
            "4. Daily management safety walks – Siemens Energy PM + Subcontractor PM.",
            "5. Consider partial termination / back-charge for additional EHS supervision if no improvement in 7 days."
        ],
        "visual": "Trend chart placeholder: Bar chart showing rising PPE violations (Week 1-4) in RED. Use repetition icon (circular arrows) in AMBER. Risk badge RED_CRITICAL flashing effect suggestion (PowerPoint animation: pulse). Bottom quote box: 'Repeated violation = Systemic Failure' – TEAL_ACCENT."
    },
    {
        "type": "action",
        "title": "Action Plan & Recommendations – Immediate Management Intervention Required",
        "visual": "Executive summary slide: DEEP_PETROL background with 3 columns in WHITE cards. Column 1: IMMEDIATE (0-24h) in RED header. Column 2: SHORT-TERM (1-7 days) in AMBER header. Column 3: STRATEGIC (7-30 days) in TEAL header. Bottom banner: Management Decision Matrix. Use icons: stop, training, contract, audit."
    }
]

# Helper to create issue slide
def create_issue_slide(slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank
    set_slide_bg(slide, DEEP_PETROL)
    
    # Top header bar
    header = add_shape(slide, Inches(0), Inches(0), Inches(13.33), Inches(1.1), fill_color=DARK_PETROL)
    # Teal accent line below header
    add_shape(slide, Inches(0), Inches(1.1), Inches(13.33), Pt(4), fill_color=TEAL_ACCENT)
    # Left red critical vertical bar
    if slide_data["risk_level"] == "CRITICAL":
        add_shape(slide, Inches(0), Inches(1.1), Pt(8), Inches(6.4), fill_color=RED_CRITICAL)
    else:
        add_shape(slide, Inches(0), Inches(1.1), Pt(8), Inches(6.4), fill_color=AMBER)

    # Slide number
    add_text_box(slide, Inches(0.3), Inches(0.15), Inches(1), Inches(0.4), f"OBS {slide_data['number']}", font_size=12, bold=True, color=TEAL_LIGHT)
    # Title
    add_text_box(slide, Inches(0.3), Inches(0.35), Inches(9.5), Inches(0.7), slide_data["title"], font_size=22, bold=True, color=WHITE)

    # Risk badge top right
    risk_color = RED_CRITICAL if slide_data["risk_level"] == "CRITICAL" else AMBER
    risk_text_color = WHITE if slide_data["risk_level"] == "CRITICAL" else DARK_PETROL
    badge = add_shape(slide, Inches(10.2), Inches(0.2), Inches(2.8), Inches(0.7), fill_color=risk_color)
    add_text_box(slide, Inches(10.2), Inches(0.2), Inches(2.8), Inches(0.7), f"RISK: {slide_data['risk_level']}", font_size=16, bold=True, color=risk_text_color, alignment=PP_ALIGN.CENTER)

    # Non-compliance box
    nc_box = add_shape(slide, Inches(0.3), Inches(1.4), Inches(4.2), Inches(4.8), fill_color=RGBColor(0x14, 0x3D, 0x4E))
    nc_box.line.color.rgb = TEAL_ACCENT
    nc_box.line.width = Pt(1)
    add_text_box(slide, Inches(0.45), Inches(1.45), Inches(3.9), Inches(0.4), "NON-COMPLIANCE DETAILS", font_size=12, bold=True, color=TEAL_LIGHT)
    tf_box = slide.shapes.add_textbox(Inches(0.45), Inches(1.85), Inches(3.9), Inches(4.3))
    tf = tf_box.text_frame
    tf.word_wrap = True
    for idx, line in enumerate(slide_data["non_compliance"]):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(11)
        p.font.color.rgb = WHITE
        p.font.name = "Calibri"
        p.space_after = Pt(8)
        p.level = 0

    # Risk description box (middle)
    risk_box = add_shape(slide, Inches(4.7), Inches(1.4), Inches(3.9), Inches(1.9), fill_color=RGBColor(0x1A, 0x1A, 0x1A) if slide_data["risk_level"]=="CRITICAL" else RGBColor(0x2B, 0x20, 0x00))
    # add border
    risk_box.line.color.rgb = risk_color
    risk_box.line.width = Pt(2)
    add_text_box(slide, Inches(4.85), Inches(1.45), Inches(3.6), Inches(0.3), "ASSOCIATED RISK & IMPACT", font_size=11, bold=True, color=risk_color)
    add_text_box(slide, Inches(4.85), Inches(1.75), Inches(3.6), Inches(1.4), slide_data["risk_desc"], font_size=10.5, bold=False, color=LIGHT_GREY)

    # Corrective Action box
    ca_box = add_shape(slide, Inches(4.7), Inches(3.5), Inches(3.9), Inches(2.8), fill_color=WHITE)
    add_text_box(slide, Inches(4.85), Inches(3.55), Inches(3.6), Inches(0.3), "REQUIRED CORRECTIVE ACTION", font_size=11, bold=True, color=DARK_PETROL)
    tf2 = slide.shapes.add_textbox(Inches(4.85), Inches(3.9), Inches(3.6), Inches(2.3)).text_frame
    tf2.word_wrap = True
    for idx, line in enumerate(slide_data["corrective"]):
        if idx == 0:
            p = tf2.paragraphs[0]
        else:
            p = tf2.add_paragraph()
        p.text = line
        p.font.size = Pt(10.5)
        p.font.color.rgb = DARK_PETROL
        p.font.name = "Calibri"
        p.space_after = Pt(5)

    # Visual & Photo placeholder
    photo_box = add_shape(slide, Inches(8.9), Inches(1.4), Inches(4.1), Inches(3.2), fill_color=RGBColor(0x0A, 0x22, 0x2B))
    photo_box.line.color.rgb = LIGHT_GREY2
    photo_box.line.width = Pt(1)
    # dashed look simulated with label
    add_text_box(slide, Inches(9.0), Inches(1.5), Inches(3.9), Inches(0.3), "📷 EVIDENCE / REFERENCE", font_size=10, bold=True, color=LIGHT_GREY2)
    add_text_box(slide, Inches(9.0), Inches(1.9), Inches(3.9), Inches(2.0), "[ Insert Site Photograph\n  Fig "+slide_data["number"]+".1 – Non-Compliance Evidence ]\n\nStandard Ref:\n• ISO 45001 / ISO 14001\n• SE EHS Contractor Req.\n• OSHA / Local Regs", font_size=9, bold=False, color=LIGHT_GREY2, alignment=PP_ALIGN.CENTER)

    # Visual suggestion box bottom right
    vs_box = add_shape(slide, Inches(8.9), Inches(4.8), Inches(4.1), Inches(1.5), fill_color=RGBColor(0x0E, 0x3A, 0x4A))
    add_text_box(slide, Inches(9.0), Inches(4.85), Inches(3.9), Inches(0.25), "VISUAL & DESIGN SUGGESTION", font_size=9, bold=True, color=TEAL_LIGHT)
    add_text_box(slide, Inches(9.0), Inches(5.1), Inches(3.9), Inches(1.1), slide_data["visual"], font_size=8, bold=False, color=LIGHT_GREY)

    # Footer
    add_shape(slide, Inches(0), Inches(6.5), Inches(13.33), Pt(1), fill_color=LIGHT_GREY2)
    add_text_box(slide, Inches(0.3), Inches(6.6), Inches(6), Inches(0.3), "Siemens Energy | Zero Harm Culture | Confidential", font_size=8, color=LIGHT_GREY2)
    add_text_box(slide, Inches(10), Inches(6.6), Inches(3), Inches(0.3), f"Slide {slide_data['number']} | SE-EHS-NC-2026", font_size=8, color=LIGHT_GREY2, alignment=PP_ALIGN.RIGHT)

# Build slides
for data in slides_data:
    if data["type"] == "title":
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_slide_bg(slide, DEEP_PETROL)
        # Accent shapes
        add_shape(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.15), fill_color=TEAL_ACCENT)
        add_shape(slide, Inches(0.8), Inches(1.0), Inches(3), Pt(3), fill_color=AMBER)
        add_text_box(slide, Inches(0.8), Inches(0.3), Inches(3), Inches(0.4), "SIEMENS ENERGY", font_size=14, bold=True, color=TEAL_LIGHT)
        add_text_box(slide, Inches(9.5), Inches(0.3), Inches(3.5), Inches(0.4), "CONFIDENTIAL", font_size=10, bold=True, color=AMBER, alignment=PP_ALIGN.RIGHT)
        
        add_text_box(slide, Inches(0.8), Inches(1.4), Inches(11.5), Inches(1.8), data["title"], font_size=32, bold=True, color=WHITE)
        add_text_box(slide, Inches(0.8), Inches(3.2), Inches(8), Inches(0.6), data["subtitle"], font_size=18, bold=False, color=TEAL_LIGHT)
        add_shape(slide, Inches(0.8), Inches(3.9), Inches(11.5), Pt(1), fill_color=LIGHT_GREY2)
        add_text_box(slide, Inches(0.8), Inches(4.2), Inches(6), Inches(2), data["meta"], font_size=11, color=LIGHT_GREY)
        
        # Right side graphic
        add_shape(slide, Inches(8.5), Inches(4.2), Inches(4), Inches(2.2), fill_color=RGBColor(0x14, 0x3D, 0x4E))
        add_text_box(slide, Inches(8.7), Inches(4.3), Inches(3.6), Inches(0.3), "VISUAL & DESIGN SUGGESTION", font_size=10, bold=True, color=TEAL_LIGHT)
        add_text_box(slide, Inches(8.7), Inches(4.6), Inches(3.6), Inches(1.6), data["visual"], font_size=9, color=LIGHT_GREY)
        
        add_text_box(slide, Inches(0.8), Inches(6.8), Inches(11), Inches(0.3), "Zero Harm | Safety is Non-Negotiable | Prepared for Top Management Review & Decision", font_size=9, bold=True, color=AMBER, alignment=PP_ALIGN.CENTER)

    elif data["type"] == "issue":
        create_issue_slide(data)
    elif data["type"] == "action":
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_slide_bg(slide, DEEP_PETROL)
        add_shape(slide, Inches(0), Inches(0), Inches(13.33), Inches(1.0), fill_color=DARK_PETROL)
        add_shape(slide, Inches(0), Inches(1.0), Inches(13.33), Pt(4), fill_color=TEAL_ACCENT)
        add_text_box(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.7), "ACTION PLAN & RECOMMENDATIONS – IMMEDIATE MANAGEMENT INTERVENTION REQUIRED", font_size=20, bold=True, color=WHITE)
        
        # 3 columns
        # Immediate
        col1 = add_shape(slide, Inches(0.3), Inches(1.4), Inches(4.1), Inches(4.5), fill_color=WHITE)
        add_shape(slide, Inches(0.3), Inches(1.4), Inches(4.1), Inches(0.5), fill_color=RED_CRITICAL)
        add_text_box(slide, Inches(0.4), Inches(1.4), Inches(3.9), Inches(0.5), "IMMEDIATE (0-24 HOURS)", font_size=12, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
        immediate_text = """• STOP WORK Order for all critical violations (PTW, Excavation, Diesel Storage, Leaking Equipment)

• Evacuate & secure hazardous areas – hard barricade + fire watch

• Remove all non-inducted personnel – 100% gate verification

• Ground all leaking equipment – Tag Out

• Issue Formal NCR Level 1 & 2 and contractual warning letter

• Convene Emergency EHS Meeting – Subcontractor MD/PM to attend

• Decision Required: Management approval for stop-work enforcement"""
        add_text_box(slide, Inches(0.4), Inches(2.0), Inches(3.9), Inches(3.8), immediate_text, font_size=10, color=DARK_PETROL)

        # Short term
        col2 = add_shape(slide, Inches(4.6), Inches(1.4), Inches(4.1), Inches(4.5), fill_color=WHITE)
        add_shape(slide, Inches(4.6), Inches(1.4), Inches(4.1), Inches(0.5), fill_color=AMBER)
        add_text_box(slide, Inches(4.7), Inches(1.4), Inches(3.9), Inches(0.5), "SHORT-TERM (1-7 DAYS)", font_size=12, bold=True, color=DARK_PETROL, alignment=PP_ALIGN.CENTER)
        short_text = """• Full PTW System Audit & Re-Training – all supervisors certified

• Compliant diesel storage – 110% bund, fire extinguishers, signage, earthing

• Reinstate traffic signage – retro-reflective + pedestrian segregation + flagmen

• Waste management system – color-coded skips, labeling, training, focal point

• Excavation compliance – hard barricades, access ladders, competent person checks

• PPE Campaign & BBS Launch – daily checks, scorecard, incentives

• EHS Directives Log – daily sign-off by Subcontractor PM

• Decision Required: Approve cost for additional SE EHS supervision & back-charge"""
        add_text_box(slide, Inches(4.7), Inches(2.0), Inches(3.9), Inches(3.8), short_text, font_size=10, color=DARK_PETROL)

        # Strategic
        col3 = add_shape(slide, Inches(8.9), Inches(1.4), Inches(4.1), Inches(4.5), fill_color=WHITE)
        add_shape(slide, Inches(8.9), Inches(1.4), Inches(4.1), Inches(0.5), fill_color=TEAL_ACCENT)
        add_text_box(slide, Inches(9.0), Inches(1.4), Inches(3.9), Inches(0.5), "STRATEGIC (7-30 DAYS) – CULTURE FIX", font_size=11, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
        strat_text = """• Contractual Enforcement: Implement penalty matrix linked to EHS KPIs (TRIR, NCR closure, PTW compliance)

• Leadership Accountability: Replace non-performing supervisors – competency assessment required

• System Upgrade: QR/Biometric induction tracking + digital PTW system + equipment checklist app

• KPI & Scorecard: Weekly EHS compliance dashboard to Top Management – Red/Amber/Green rating

• Independent Audit: 3rd party EHS audit of subcontractor management system

• Zero Harm Commitment: Re-sign Life Saving Rules – Subcontractor Top Management

• Contingency: Prepare demobilization / replacement plan if no sustained improvement

• Decision Required: Authorize contractual action up to partial termination if critical violations recur"""
        add_text_box(slide, Inches(9.0), Inches(2.0), Inches(3.9), Inches(3.8), strat_text, font_size=10, color=DARK_PETROL)

        # Bottom decision box
        add_shape(slide, Inches(0.3), Inches(6.1), Inches(12.7), Inches(1.1), fill_color=RGBColor(0x14, 0x3D, 0x4E))
        add_text_box(slide, Inches(0.4), Inches(6.15), Inches(12.5), Inches(0.25), "MANAGEMENT DECISION MATRIX – REQUIRED TODAY", font_size=11, bold=True, color=AMBER)
        decision_text = "1. Approve IMMEDIATE Stop-Work Authority to EHS for Critical Violations  |  2. Authorize Formal Contractual Warning & Penalty Enforcement  |  3. Mandate Subcontractor Top Management Presence On-Site within 48h  |  4. Approve Budget for Additional EHS Controls (Bunded Storage, Signage, Barricades) – Back-charge to Subcontractor  |  5. Endorse Zero Tolerance Policy: 3 Strikes = Removal – Communicate in Townhall"
        add_text_box(slide, Inches(0.4), Inches(6.4), Inches(12.5), Inches(0.7), decision_text, font_size=10, bold=False, color=WHITE)

        # Visual suggestion footer
        add_text_box(slide, Inches(0.3), Inches(7.25), Inches(12.7), Inches(0.2), "Visual: Use Siemens Energy corporate template – Deep Petrol #0E2F3E background, White cards for readability, Red for Immediate, Amber for Short-Term, Teal for Strategic. Add icons: Stop Hand, Training, Contract, Audit. Animate columns sequentially. End with 'Zero Harm is Non-Negotiable' tagline in Teal Light.", font_size=7.5, color=LIGHT_GREY2)

output_path = "/home/user/ahmed/Subcontractor_EHS_Non-Compliance_Report_Siemens_Energy.pptx"
prs.save(output_path)
print(f"Saved to {output_path}")
