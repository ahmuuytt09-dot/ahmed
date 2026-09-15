#!/usr/bin/env python3
"""Generate professional corporate EHS Site Inspection PowerPoint.

KAZ Power Plant Upgrade Project | EHS | 15 September 2026
7 slides, 16:9, navy / white / light-grey corporate theme.
Photo placeholders only — no stock or AI imagery.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE

# ---------------------------------------------------------------- palette
NAVY        = RGBColor(0x0F, 0x2A, 0x44)
NAVY_DARK   = RGBColor(0x0A, 0x1E, 0x32)
NAVY_PANEL  = RGBColor(0x14, 0x36, 0x58)
NAVY_LINE   = RGBColor(0x1E, 0x4A, 0x73)
ACCENT      = RGBColor(0x2E, 0x86, 0xAB)
ACCENT_LT   = RGBColor(0x7F, 0xB4, 0xD1)
ACCENT_BG   = RGBColor(0xE9, 0xF1, 0xF7)
GREY_BG     = RGBColor(0xF2, 0xF4, 0xF7)
GREY_MID    = RGBColor(0xE3, 0xE8, 0xEF)
GREY_BORDER = RGBColor(0xCF, 0xD6, 0xE0)
GREY_DASH   = RGBColor(0x9A, 0xA8, 0xBC)
TEXT        = RGBColor(0x1B, 0x2A, 0x41)
MUTED       = RGBColor(0x5C, 0x6B, 0x80)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
AMBER       = RGBColor(0xB7, 0x79, 0x1F)
AMBER_BG    = RGBColor(0xFE, 0xF6, 0xE9)
AMBER_LINE  = RGBColor(0xE0, 0xC4, 0x8C)
COVER_SUB   = RGBColor(0xC4, 0xD2, 0xE2)

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
FOOTER_H = 0.42
HEADER_H = 1.12
FONT = "Calibri"

FOOTER_TEXT = "KAZ Power Plant Upgrade Project  |  EHS  |  15 September 2026"

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
BLANK = prs.slide_layouts[6]

# ---------------------------------------------------------------- helpers
def set_bg(slide, rgb):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = rgb

def no_line(shape):
    shape.line.fill.background()

def style_para(p, text, size=10, bold=False, italic=False, color=TEXT,
               align=PP_ALIGN.LEFT, space_after=0, space_before=0,
               line_spacing=1.08, font_name=FONT):
    p.text = ""
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font_name
    return p

def add_para(tf, text, size=10, bold=False, italic=False, color=TEXT,
             align=PP_ALIGN.LEFT, space_after=0, space_before=0,
             line_spacing=1.08):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    return p

def textbox(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.text_frame.word_wrap = wrap
    tb.text_frame.vertical_anchor = anchor
    tb.text_frame.margin_left = Inches(0.02)
    tb.text_frame.margin_right = Inches(0.02)
    tb.text_frame.margin_top = Inches(0.02)
    tb.text_frame.margin_bottom = Inches(0.02)
    no_line(tb)
    return tb

def rect(slide, l, t, w, h, fill=None, line=None, line_w=1.0,
         dash=None, rounded=False, radius=0.08, name=None):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shp_type, Inches(l), Inches(t), Inches(w), Inches(h))
    if rounded:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        no_line(sp)
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
        if dash is not None:
            sp.line.dash_style = dash
    try:
        sp.shadow.inherit = False
    except Exception:
        pass
    if name:
        try:
            sp.name = name
        except Exception:
            pass
    return sp

def oval(slide, l, t, size, fill=None, line=None, line_w=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(l), Inches(t), Inches(size), Inches(size))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        no_line(sp)
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    try:
        sp.shadow.inherit = False
    except Exception:
        pass
    return sp

def icon_badge(slide, l, t, size, glyph, fill_rgb, glyph_rgb, glyph_pt=None):
    b = oval(slide, l, t, size, fill=fill_rgb)
    tf = b.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.01); tf.margin_right = Inches(0.01)
    tf.margin_top = Inches(0.01); tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = glyph
    run.font.size = Pt(glyph_pt or max(10, size * 34))
    run.font.bold = True
    run.font.color.rgb = glyph_rgb
    run.font.name = FONT
    return b

def add_footer(slide, idx, total=7, on_cover=False):
    # thin separator on cover (cover bg is navy)
    if on_cover:
        rect(slide, 0, SLIDE_H - FOOTER_H - 0.012, SLIDE_W, 0.012, fill=NAVY_LINE)
    bar = rect(slide, 0, SLIDE_H - FOOTER_H, SLIDE_W, FOOTER_H, fill=NAVY)
    # left footer text
    tb = textbox(slide, MARGIN, SLIDE_H - FOOTER_H + 0.09, 9.5, 0.26)
    style_para(tb.text_frame.paragraphs[0], FOOTER_TEXT, size=8, color=WHITE, align=PP_ALIGN.LEFT)
    # right page number
    tb2 = textbox(slide, SLIDE_W - MARGIN - 1.8, SLIDE_H - FOOTER_H + 0.09, 1.8, 0.26)
    style_para(tb2.text_frame.paragraphs[0], f"{idx:02d}  /  {total:02d}", size=8, color=COVER_SUB, align=PP_ALIGN.RIGHT)
    return bar

def add_header(slide, number, eyebrow, title):
    rect(slide, 0, 0, SLIDE_W, HEADER_H, fill=NAVY)
    rect(slide, 0, HEADER_H, SLIDE_W, 0.045, fill=ACCENT)
    # eyebrow
    tb = textbox(slide, MARGIN, 0.14, 9.6, 0.24)
    style_para(tb.text_frame.paragraphs[0], eyebrow, size=8, bold=True, color=ACCENT_LT, align=PP_ALIGN.LEFT)
    # title
    tb2 = textbox(slide, MARGIN, 0.36, 9.6, 0.6)
    style_para(tb2.text_frame.paragraphs[0], title, size=24, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
    # slide number right
    tb3 = textbox(slide, SLIDE_W - MARGIN - 1.6, 0.12, 1.6, 0.55)
    style_para(tb3.text_frame.paragraphs[0], f"{number:02d}", size=30, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    tb4 = textbox(slide, SLIDE_W - MARGIN - 1.6, 0.62, 1.6, 0.3)
    style_para(tb4.text_frame.paragraphs[0], "EHS  •  FORMAL", size=7.5, bold=True, color=COVER_SUB, align=PP_ALIGN.RIGHT)

def photo_placeholder(slide, l, t, w, h, label, hint="Delete this box and insert the photo at the same size"):
    """Professional dashed photo placeholder box + top tag. Returns outer shape."""
    box = rect(slide, l, t, w, h, fill=GREY_BG, line=GREY_DASH,
               line_w=1.4, dash=MSO_LINE_DASH_STYLE.DASH, rounded=True,
               radius=0.05, name=label)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.12); tf.margin_bottom = Inches(0.12)
    # icon glyph
    style_para(tf.paragraphs[0], "\u25EB", size=30, color=GREY_DASH, align=PP_ALIGN.CENTER, space_after=4)
    add_para(tf, label, size=11, bold=True, color=TEXT, align=PP_ALIGN.CENTER, space_after=3)
    add_para(tf, hint, size=7.5, italic=True, color=MUTED, align=PP_ALIGN.CENTER)
    # top tag pill
    tag_w, tag_h = 1.75, 0.27
    tag = rect(slide, l + w/2 - tag_w/2, t - tag_h/2 + 0.02, tag_w, tag_h,
               fill=NAVY, rounded=True, radius=0.5)
    tf2 = tag.text_frame
    tf2.word_wrap = False
    tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf2.margin_left = Inches(0.05); tf2.margin_right = Inches(0.05)
    tf2.margin_top = Inches(0.0); tf2.margin_bottom = Inches(0.0)
    style_para(tf2.paragraphs[0], "PHOTO PLACEHOLDER", size=7, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    return box

# ================================================================ SLIDE 1 — COVER
s1 = prs.slides.add_slide(BLANK)
set_bg(s1, NAVY)
# right visual panel
rect(s1, 8.35, 0, SLIDE_W - 8.35, SLIDE_H - FOOTER_H, fill=NAVY_PANEL)
rect(s1, 8.35, 0, 0.045, SLIDE_H - FOOTER_H, fill=ACCENT)
# subtle horizontal grid lines in panel
for i, y in enumerate([1.35, 2.75, 4.15, 5.30]):
    rect(s1, 8.75, y, SLIDE_W - 8.75 - 0.6, 0.012, fill=NAVY_LINE)
# emblem: outer ring + inner disc + bolt
ring = oval(s1, 9.78, 1.55, 1.95, fill=None, line=ACCENT, line_w=1.6)
disc = oval(s1, 10.02, 1.79, 1.47, fill=NAVY_DARK, line=NAVY_LINE, line_w=1.0)
tbd = textbox(s1, 10.02, 2.02, 1.47, 0.9, anchor=MSO_ANCHOR.MIDDLE)
style_para(tbd.text_frame.paragraphs[0], "\u26A1", size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
tbe = textbox(s1, 8.75, 3.72, SLIDE_W - 8.75 - 0.6, 0.28)
style_para(tbe.text_frame.paragraphs[0], "ENERGIZED FACILITY", size=9, bold=True, color=ACCENT_LT, align=PP_ALIGN.CENTER)
tbe2 = textbox(s1, 8.75, 3.98, SLIDE_W - 8.75 - 0.6, 0.28)
style_para(tbe2.text_frame.paragraphs[0], "Strict EHS compliance required", size=9, italic=True, color=COVER_SUB, align=PP_ALIGN.CENTER)
# three mini pills in panel
pill_labels = [("W", "Work Permit"), ("I", "Induction"), ("P", "PPE")]
py = 4.62
for glyph, lab in pill_labels:
    pill = rect(s1, 9.30, py, 2.55, 0.42, fill=NAVY_DARK, line=NAVY_LINE, line_w=0.9, rounded=True, radius=0.35)
    # glyph
    tbp = textbox(s1, 9.38, py + 0.045, 0.4, 0.33, anchor=MSO_ANCHOR.MIDDLE)
    style_para(tbp.text_frame.paragraphs[0], glyph, size=11, bold=True, color=ACCENT_LT, align=PP_ALIGN.CENTER)
    rect(s1, 9.80, py + 0.09, 0.012, 0.24, fill=NAVY_LINE)
    tbp2 = textbox(s1, 9.90, py + 0.045, 1.85, 0.33, anchor=MSO_ANCHOR.MIDDLE)
    style_para(tbp2.text_frame.paragraphs[0], lab, size=9.5, bold=False, color=WHITE, align=PP_ALIGN.LEFT)
    py += 0.56

# left content
badge = rect(s1, MARGIN, 0.62, 2.55, 0.36, fill=NAVY_DARK, line=NAVY_LINE, line_w=0.9, rounded=True, radius=0.4)
tfb = badge.text_frame
tfb.word_wrap = False
tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
tfb.margin_left = Inches(0.12); tfb.margin_right = Inches(0.12)
style_para(tfb.paragraphs[0], "EHS  \u2022  SITE INSPECTION", size=8.5, bold=True, color=ACCENT_LT, align=PP_ALIGN.CENTER)

tb = textbox(s1, MARGIN, 1.25, 7.0, 1.6)
style_para(tb.text_frame.paragraphs[0], "EHS Site Inspection \u2013 Non-Compliance", size=38, bold=True, color=WHITE, align=PP_ALIGN.LEFT, line_spacing=1.0)
tb2 = textbox(s1, MARGIN, 2.95, 7.0, 0.55)
style_para(tb2.text_frame.paragraphs[0], "Energized Power Plant", size=20, color=COVER_SUB, align=PP_ALIGN.LEFT)
# accent divider
rect(s1, MARGIN, 3.62, 0.95, 0.055, fill=ACCENT)
rect(s1, MARGIN + 1.02, 3.62, 2.6, 0.055, fill=NAVY_LINE)
# metadata rows
meta = [("Date", "15 September 2026"), ("Prepared by", "Ahmed Al-Mansoury"), ("Title", "EHS Manager")]
my = 4.05
for k, v in meta:
    tmk = textbox(s1, MARGIN, my, 1.55, 0.3)
    style_para(tmk.text_frame.paragraphs[0], k.upper(), size=8, bold=True, color=ACCENT_LT, align=PP_ALIGN.LEFT)
    tmv = textbox(s1, MARGIN + 1.55, my - 0.02, 4.6, 0.32)
    style_para(tmv.text_frame.paragraphs[0], v, size=12, color=WHITE, align=PP_ALIGN.LEFT)
    my += 0.42
# project strip
rect(s1, MARGIN, 5.62, 7.0, 0.014, fill=NAVY_LINE)
tb3 = textbox(s1, MARGIN, 5.82, 7.0, 0.5)
style_para(tb3.text_frame.paragraphs[0], "KAZ Power Plant Upgrade Project  \u00b7  Formal EHS communication to the subcontractor",
           size=9.5, italic=True, color=COVER_SUB, align=PP_ALIGN.LEFT)
# small classification bottom-left above footer
tb4 = textbox(s1, MARGIN, SLIDE_H - FOOTER_H - 0.55, 7.0, 0.3)
style_para(tb4.text_frame.paragraphs[0], "Document type:  EHS Inspection Findings & Formal Warning", size=8, color=MUTED if False else COVER_SUB, align=PP_ALIGN.LEFT)
add_footer(s1, 1, on_cover=True)

# ================================================================ SLIDE 2 — FINDINGS
s2 = prs.slides.add_slide(BLANK)
set_bg(s2, WHITE)
add_header(s2, 2, "KAZ POWER PLANT UPGRADE PROJECT  \u2022  EHS SITE INSPECTION", "Observed EHS Non-Compliances")
findings = [
    ("\u2630", "FINDING 01  \u2014  WORK PERMIT", "Workers performing activities without valid Work Permits."),
    ("\u25C9", "FINDING 02  \u2014  SITE INDUCTION", "Workers working without completing the required site induction."),
    ("\u2B22", "FINDING 03  \u2014  PPE", "Workers not wearing the required PPE."),
    ("\u26A1", "FINDING 04  \u2014  EHS CONTROLS", "Activities being performed without the required EHS controls."),
]
card_w = (SLIDE_W - MARGIN*2 - 0.35) / 2
card_h = 1.52
top0 = 1.62
for i, (glyph, eyebrow, desc) in enumerate(findings):
    r, c = divmod(i, 2)
    l = MARGIN + c * (card_w + 0.35)
    t = top0 + r * (card_h + 0.28)
    card = rect(s2, l, t, card_w, card_h, fill=WHITE, line=GREY_BORDER, line_w=0.9, rounded=True, radius=0.06)
    rect(s2, l, t, 0.09, card_h, fill=ACCENT)  # left accent (slightly overlaps rounded corner; acceptable, clean)
    icon_badge(s2, l + 0.32, t + 0.40, 0.68, glyph, NAVY, WHITE, glyph_pt=17)
    t1 = textbox(s2, l + 1.20, t + 0.24, card_w - 1.45, 0.30)
    style_para(t1.text_frame.paragraphs[0], eyebrow, size=8.5, bold=True, color=ACCENT, align=PP_ALIGN.LEFT)
    t2 = textbox(s2, l + 1.20, t + 0.55, card_w - 1.45, 0.75, anchor=MSO_ANCHOR.TOP)
    style_para(t2.text_frame.paragraphs[0], desc, size=12.5, color=TEXT, align=PP_ALIGN.LEFT, line_spacing=1.12)
# highlighted note
note_t = top0 + 2 * card_h + 2 * 0.28 + 0.10
note_h = 1.22
note = rect(s2, MARGIN, note_t, SLIDE_W - MARGIN*2, note_h, fill=ACCENT_BG, rounded=True, radius=0.08)
rect(s2, MARGIN, note_t, 0.10, note_h, fill=NAVY)
icon_badge(s2, MARGIN + 0.35, note_t + 0.30, 0.60, "\u26A0", NAVY, WHITE, glyph_pt=15)
# NOTE label
tb = textbox(s2, MARGIN + 1.12, note_t + 0.16, SLIDE_W - MARGIN*2 - 1.35, 0.26)
style_para(tb.text_frame.paragraphs[0], "NOTE  \u2014  ENERGIZED FACILITY", size=8.5, bold=True, color=NAVY, align=PP_ALIGN.LEFT)
tb2 = textbox(s2, MARGIN + 1.12, note_t + 0.42, SLIDE_W - MARGIN*2 - 1.35, 0.66)
style_para(tb2.text_frame.paragraphs[0],
           "\u201CThe activities were observed within an energized power plant; therefore, strict compliance with the site EHS requirements is mandatory.\u201D",
           size=11.5, italic=True, color=TEXT, align=PP_ALIGN.LEFT, line_spacing=1.15)
add_footer(s2, 2)

# ================================================================ SLIDE 3 — PERMIT & INDUCTION
s3 = prs.slides.add_slide(BLANK)
set_bg(s3, WHITE)
add_header(s3, 3, "KAZ POWER PLANT UPGRADE PROJECT  \u2022  EHS SITE INSPECTION", "Work Permit & Site Induction Non-Compliance")
intro = textbox(s3, MARGIN, 1.38, SLIDE_W - MARGIN*2, 0.30)
style_para(intro.text_frame.paragraphs[0], "Insert the corresponding inspection photographs in the placeholders below. Captions describe each finding.",
           size=10, italic=True, color=MUTED, align=PP_ALIGN.LEFT)
ph_w = (SLIDE_W - MARGIN*2 - 0.5) / 2
ph_h = 3.70
ph_t = 1.95
photo_placeholder(s3, MARGIN, ph_t, ph_w, ph_h, "Insert Work Permit Non-Compliance Photo")
photo_placeholder(s3, MARGIN + ph_w + 0.5, ph_t, ph_w, ph_h, "Insert Site Induction Non-Compliance Photo")
caps = ["Work Permit non-compliance observed during inspection.", "Site induction non-compliance observed during inspection."]
for j, cap in enumerate(caps):
    l = MARGIN + j * (ph_w + 0.5)
    # caption bar
    cb = rect(s3, l, ph_t + ph_h + 0.16, ph_w, 0.52, fill=GREY_BG, rounded=True, radius=0.15)
    tfc = cb.text_frame
    tfc.word_wrap = True
    tfc.vertical_anchor = MSO_ANCHOR.MIDDLE
    tfc.margin_left = Inches(0.18); tfc.margin_right = Inches(0.18)
    style_para(tfc.paragraphs[0], "Caption:  " + cap, size=9.5, color=TEXT, align=PP_ALIGN.LEFT)
    # small label above caption? finding tag
    tag = textbox(s3, l, ph_t + ph_h + 0.68, ph_w, 0.24)
    style_para(tag.text_frame.paragraphs[0], f"FINDING 0{j+1}  \u2022  PHOTO EVIDENCE TO BE INSERTED", size=7.5, bold=True, color=MUTED, align=PP_ALIGN.LEFT)
add_footer(s3, 3)

# ================================================================ SLIDE 4 — PPE
s4 = prs.slides.add_slide(BLANK)
set_bg(s4, WHITE)
add_header(s4, 4, "KAZ POWER PLANT UPGRADE PROJECT  \u2022  EHS SITE INSPECTION", "PPE Non-Compliance")
intro = textbox(s4, MARGIN, 1.38, SLIDE_W - MARGIN*2, 0.30)
style_para(intro.text_frame.paragraphs[0], "Insert the actual PPE inspection photographs in the placeholders below.",
           size=10, italic=True, color=MUTED, align=PP_ALIGN.LEFT)
n = 3
gap = 0.30
ph_w = (SLIDE_W - MARGIN*2 - gap*(n-1)) / n
ph_h = 3.35
ph_t = 1.95
tags = ["PHOTO A", "PHOTO B", "PHOTO C"]
for j in range(n):
    l = MARGIN + j * (ph_w + gap)
    photo_placeholder(s4, l, ph_t, ph_w, ph_h, "Insert PPE Non-Compliance Photo")
    tg = textbox(s4, l, ph_t + ph_h + 0.10, ph_w, 0.24)
    style_para(tg.text_frame.paragraphs[0], f"{tags[j]}  \u2022  PPE FINDING", size=7.5, bold=True, color=MUTED, align=PP_ALIGN.LEFT)
# caption bar full width
cap = rect(s4, MARGIN, ph_t + ph_h + 0.42, SLIDE_W - MARGIN*2, 0.58, fill=ACCENT_BG, rounded=True, radius=0.12)
rect(s4, MARGIN, ph_t + ph_h + 0.42, 0.10, 0.58, fill=NAVY)
tfc = cap.text_frame
tfc.word_wrap = True
tfc.vertical_anchor = MSO_ANCHOR.MIDDLE
tfc.margin_left = Inches(0.35); tfc.margin_right = Inches(0.2)
style_para(tfc.paragraphs[0], "Personnel were observed without the required PPE.", size=11, italic=True, color=TEXT, align=PP_ALIGN.LEFT)
add_footer(s4, 4)

# ================================================================ SLIDE 5 — ACTIVITIES
s5 = prs.slides.add_slide(BLANK)
set_bg(s5, WHITE)
add_header(s5, 5, "KAZ POWER PLANT UPGRADE PROJECT  \u2022  EHS SITE INSPECTION", "Activities Observed During Inspection")
intro = textbox(s5, MARGIN, 1.38, SLIDE_W - MARGIN*2, 0.30)
style_para(intro.text_frame.paragraphs[0], "Activities observed during the inspection. Insert one photograph per activity.",
           size=10, italic=True, color=MUTED, align=PP_ALIGN.LEFT)
acts = [
    ("01  \u2014  CEMENT WORKS", "Cement works observed during inspection."),
    ("02  \u2014  FENCE INSTALLATION", "Fence installation observed during inspection."),
    ("03  \u2014  TRANSFORMER RAIL PAINTING", "Transformer rail painting observed during inspection."),
]
n = 3
gap = 0.30
col_w = (SLIDE_W - MARGIN*2 - gap*(n-1)) / n
col_t = 1.85
head_h = 0.52
ph_h5 = 2.95
desc_h = 0.60
for j, (ahead, adesc) in enumerate(acts):
    l = MARGIN + j * (col_w + gap)
    # header bar
    hb = rect(s5, l, col_t, col_w, head_h, fill=NAVY, rounded=True, radius=0.22)
    tfh = hb.text_frame
    tfh.word_wrap = True
    tfh.vertical_anchor = MSO_ANCHOR.MIDDLE
    tfh.margin_left = Inches(0.20); tfh.margin_right = Inches(0.16)
    style_para(tfh.paragraphs[0], ahead, size=9.5, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
    photo_placeholder(s5, l, col_t + head_h + 0.14, col_w, ph_h5, "Insert Inspection Photo")
    db = rect(s5, l, col_t + head_h + 0.14 + ph_h5 + 0.14, col_w, desc_h, fill=GREY_BG, rounded=True, radius=0.15)
    tfd = db.text_frame
    tfd.word_wrap = True
    tfd.vertical_anchor = MSO_ANCHOR.MIDDLE
    tfd.margin_left = Inches(0.16); tfd.margin_right = Inches(0.16)
    style_para(tfd.paragraphs[0], adesc, size=9.5, color=TEXT, align=PP_ALIGN.LEFT)
add_footer(s5, 5)

# ================================================================ SLIDE 6 — CORRECTIVE ACTIONS
s6 = prs.slides.add_slide(BLANK)
set_bg(s6, WHITE)
add_header(s6, 6, "KAZ POWER PLANT UPGRADE PROJECT  \u2022  EHS SITE INSPECTION", "Required Corrective Actions")
intro = textbox(s6, MARGIN, 1.38, SLIDE_W - MARGIN*2, 0.30)
style_para(intro.text_frame.paragraphs[0], "The following corrective actions are required before starting any activity:",
           size=10.5, italic=True, color=MUTED, align=PP_ALIGN.LEFT)
actions = [
    "Ensure valid Work Permits are obtained before starting any activity.",
    "Ensure all personnel complete the required site induction.",
    "Ensure full PPE compliance at all times.",
    "Ensure appropriate EHS controls are implemented before work starts.",
    "Ensure compliance with the requirements applicable to the energized facility.",
]
row_h = 0.80
row_gap = 0.16
row_t0 = 1.85
for i, act in enumerate(actions):
    t = row_t0 + i * (row_h + row_gap)
    row = rect(s6, MARGIN, t, SLIDE_W - MARGIN*2, row_h, fill=WHITE, line=GREY_BORDER, line_w=0.9, rounded=True, radius=0.10)
    rect(s6, MARGIN, t, 0.09, row_h, fill=ACCENT)
    icon_badge(s6, MARGIN + 0.32, t + 0.16, 0.48, str(i+1), NAVY, WHITE, glyph_pt=13)
    tb = textbox(s6, MARGIN + 1.02, t + 0.10, SLIDE_W - MARGIN*2 - 2.10, row_h - 0.2, anchor=MSO_ANCHOR.MIDDLE)
    style_para(tb.text_frame.paragraphs[0], act, size=11.5, color=TEXT, align=PP_ALIGN.LEFT, line_spacing=1.1)
    # action tag on right
    tag = rect(s6, SLIDE_W - MARGIN - 1.35, t + 0.22, 1.05, 0.36, fill=ACCENT_BG, rounded=True, radius=0.4)
    tft = tag.text_frame
    tft.word_wrap = False
    tft.vertical_anchor = MSO_ANCHOR.MIDDLE
    tft.margin_left = Inches(0.04); tft.margin_right = Inches(0.04)
    style_para(tft.paragraphs[0], "\u2713  REQUIRED", size=7.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
add_footer(s6, 6)

# ================================================================ SLIDE 7 — FORMAL WARNING
s7 = prs.slides.add_slide(BLANK)
set_bg(s7, GREY_BG)
add_header(s7, 7, "KAZ POWER PLANT UPGRADE PROJECT  \u2022  EHS SITE INSPECTION", "Formal EHS Warning")
# warning card
cw, chh = 10.9, 3.35
cl, ct = (SLIDE_W - cw)/2, 2.45
card = rect(s7, cl, ct, cw, chh, fill=WHITE, line=AMBER_LINE, line_w=1.1, rounded=True, radius=0.04)
rect(s7, cl, ct, 0.12, chh, fill=AMBER)
# badge + eyebrow
icon_badge(s7, cl + 0.50, ct + 0.38, 0.66, "\u26A0", AMBER, WHITE, glyph_pt=17)
tb = textbox(s7, cl + 1.32, ct + 0.36, cw - 2.6, 0.30)
style_para(tb.text_frame.paragraphs[0], "FORMAL EHS WARNING  \u2014  IMMEDIATE ACTION REQUIRED", size=9, bold=True, color=AMBER, align=PP_ALIGN.LEFT)
tb_to = textbox(s7, cl + cw - 2.55, ct + 0.36, 2.1, 0.30)
style_para(tb_to.text_frame.paragraphs[0], "To:  AL MIAL", size=9, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)
# recipient pill
# divider
rect(s7, cl + 0.50, ct + 1.14, cw - 1.0, 0.014, fill=GREY_MID)
# body
tb2 = textbox(s7, cl + 0.50, ct + 1.32, cw - 1.0, 1.50)
tf = tb2.text_frame
tf.word_wrap = True
style_para(tf.paragraphs[0],
           "\u201CAL MIAL is requested to take immediate corrective action and ensure full compliance with the Work Permit, induction, PPE, and applicable EHS requirements before starting any activity.",
           size=12.5, color=TEXT, align=PP_ALIGN.LEFT, space_after=8, line_spacing=1.22)
add_para(tf,
         "Any repeated non-compliance will be subject to further formal action in accordance with the project requirements.\u201D",
         size=12.5, bold=True, color=TEXT, align=PP_ALIGN.LEFT, line_spacing=1.22)
# bottom strip inside card
rect(s7, cl + 0.50, ct + chh - 0.78, cw - 1.0, 0.014, fill=GREY_MID)
tb3 = textbox(s7, cl + 0.50, ct + chh - 0.62, cw - 1.0, 0.4)
style_para(tb3.text_frame.paragraphs[0], "EHS Manager:  Ahmed Al-Mansoury   \u00b7   Date:  15 September 2026",
           size=9, color=MUTED, align=PP_ALIGN.LEFT)
add_footer(s7, 7)

# ---------------------------------------------------------------- properties + save
prs.core_properties.title = "EHS Site Inspection - Non-Compliance | Energized Power Plant"
prs.core_properties.subject = "KAZ Power Plant Upgrade Project - EHS Inspection Findings and Formal Warning"
prs.core_properties.author = "Ahmed Al-Mansoury, EHS Manager"
prs.core_properties.keywords = "EHS, inspection, non-compliance, KAZ, power plant, work permit, induction, PPE"
prs.core_properties.comments = "Photo placeholders only - insert actual site inspection photos before issue."

out = "/home/user/ahmed/KAZ_EHS_Site_Inspection_Non-Compliance_2026-09-15.pptx"
prs.save(out)
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")
